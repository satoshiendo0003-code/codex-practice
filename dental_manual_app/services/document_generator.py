"""
ドキュメント生成モジュール
Word/PDF/PowerPointの生成を担当
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image as ReportLabImage
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from typing import Dict, List
import os
from datetime import datetime


class DocumentGenerator:
    """ドキュメント生成クラス"""

    def __init__(self, output_dir: str = "output"):
        """
        Args:
            output_dir: 出力ディレクトリ
        """
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def generate_word(self, data: Dict, title: str, summary: str) -> str:
        """
        Wordドキュメントを生成

        Args:
            data: マニュアルデータ
            title: マニュアルタイトル
            summary: マニュアル概要

        Returns:
            生成したファイルのパス
        """
        doc = Document()

        # タイトル
        title_para = doc.add_heading(title, 0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 作成日
        date_para = doc.add_paragraph(f"作成日: {datetime.now().strftime('%Y年%m月%d日')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_paragraph()  # 空行

        # 概要
        doc.add_heading("概要", level=1)
        doc.add_paragraph(summary)

        # 必要な器具（データに含まれている場合）
        if "required_tools" in data and data["required_tools"]:
            doc.add_heading("必要な器具", level=1)
            for tool in data["required_tools"]:
                doc.add_paragraph(tool, style='List Bullet')

        # 推定所要時間
        if "estimated_time" in data:
            doc.add_heading("推定所要時間", level=1)
            doc.add_paragraph(data["estimated_time"])

        # 手順
        doc.add_heading("手順", level=1)

        steps = data.get("steps", [])
        for step in steps:
            # ステップタイトル
            step_title = f"ステップ {step.get('step_number', '')}: {step.get('title', '')}"
            doc.add_heading(step_title, level=2)

            # 画像を挿入（存在する場合）
            if step.get('image_path') and os.path.exists(step['image_path']):
                try:
                    doc.add_picture(step['image_path'], width=Inches(4.5))
                    # 画像の下に空行
                    doc.add_paragraph()
                except Exception as e:
                    print(f"画像の挿入に失敗: {e}")

            # 説明
            doc.add_paragraph(step.get('description', ''))

            # 重要ポイント
            if step.get('key_points'):
                doc.add_paragraph("重要ポイント:", style='List Bullet')
                for point in step['key_points']:
                    doc.add_paragraph(point, style='List Bullet 2')

            # 注意事項
            if step.get('warnings'):
                warning_para = doc.add_paragraph("⚠️ 注意事項:")
                warning_para.runs[0].font.color.rgb = RGBColor(255, 0, 0)
                for warning in step['warnings']:
                    doc.add_paragraph(warning, style='List Bullet 2')

            doc.add_paragraph()  # 空行

        # ファイルを保存
        file_path = os.path.join(self.output_dir, f"{title}.docx")
        doc.save(file_path)

        return file_path

    def generate_pdf(self, data: Dict, title: str, summary: str) -> str:
        """
        PDFドキュメントを生成

        Args:
            data: マニュアルデータ
            title: マニュアルタイトル
            summary: マニュアル概要

        Returns:
            生成したファイルのパス
        """
        file_path = os.path.join(self.output_dir, f"{title}.pdf")

        # PDFドキュメントを作成
        doc = SimpleDocTemplate(file_path, pagesize=A4)
        story = []

        # スタイルを設定
        styles = getSampleStyleSheet()

        # タイトルスタイル
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#333333'),
            spaceAfter=30,
            alignment=1  # 中央揃え
        )

        # 見出しスタイル
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#667eea'),
            spaceAfter=12,
        )

        # 本文スタイル
        body_style = ParagraphStyle(
            'CustomBody',
            parent=styles['BodyText'],
            fontSize=11,
            spaceAfter=12,
        )

        # タイトル
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 0.2*inch))

        # 作成日
        date_text = f"作成日: {datetime.now().strftime('%Y年%m月%d日')}"
        story.append(Paragraph(date_text, body_style))
        story.append(Spacer(1, 0.3*inch))

        # 概要
        story.append(Paragraph("概要", heading_style))
        story.append(Paragraph(summary, body_style))
        story.append(Spacer(1, 0.2*inch))

        # 必要な器具
        if "required_tools" in data and data["required_tools"]:
            story.append(Paragraph("必要な器具", heading_style))
            for tool in data["required_tools"]:
                story.append(Paragraph(f"• {tool}", body_style))
            story.append(Spacer(1, 0.2*inch))

        # 推定所要時間
        if "estimated_time" in data:
            story.append(Paragraph("推定所要時間", heading_style))
            story.append(Paragraph(data["estimated_time"], body_style))
            story.append(Spacer(1, 0.2*inch))

        # 手順
        story.append(Paragraph("手順", heading_style))
        story.append(Spacer(1, 0.1*inch))

        steps = data.get("steps", [])
        for step in steps:
            # ステップタイトル
            step_title = f"ステップ {step.get('step_number', '')}: {step.get('title', '')}"
            story.append(Paragraph(step_title, heading_style))

            # 画像を挿入（存在する場合）
            if step.get('image_path') and os.path.exists(step['image_path']):
                try:
                    img = ReportLabImage(step['image_path'], width=4*inch, height=3*inch)
                    story.append(Spacer(1, 0.1*inch))
                    story.append(img)
                    story.append(Spacer(1, 0.1*inch))
                except Exception as e:
                    print(f"画像の挿入に失敗: {e}")

            # 説明
            story.append(Paragraph(step.get('description', ''), body_style))

            # 重要ポイント
            if step.get('key_points'):
                story.append(Paragraph("<b>重要ポイント:</b>", body_style))
                for point in step['key_points']:
                    story.append(Paragraph(f"• {point}", body_style))

            # 注意事項
            if step.get('warnings'):
                warning_style = ParagraphStyle(
                    'Warning',
                    parent=body_style,
                    textColor=colors.red,
                )
                story.append(Paragraph("<b>⚠️ 注意事項:</b>", warning_style))
                for warning in step['warnings']:
                    story.append(Paragraph(f"• {warning}", warning_style))

            story.append(Spacer(1, 0.2*inch))

        # PDFを生成
        doc.build(story)

        return file_path
