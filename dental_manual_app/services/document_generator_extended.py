"""
拡張ドキュメント生成モジュール
PowerPoint、画像カード、テンプレート対応
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor as PptxRGBColor
from PIL import Image, ImageDraw, ImageFont
from typing import Dict, List
import os
from datetime import datetime


class ExtendedDocumentGenerator:
    """拡張ドキュメント生成クラス"""

    def __init__(self, output_dir: str = "output"):
        """
        Args:
            output_dir: 出力ディレクトリ
        """
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def generate_powerpoint(
        self,
        data: Dict,
        title: str,
        summary: str,
        template: str = "modern"
    ) -> str:
        """
        PowerPointプレゼンテーションを生成

        Args:
            data: マニュアルデータ
            title: マニュアルタイトル
            summary: マニュアル概要
            template: テンプレート名（modern, classic, simple）

        Returns:
            生成したファイルのパス
        """
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)

        # テンプレートに応じた色設定
        templates = {
            "modern": {
                "primary": PptxRGBColor(102, 126, 234),  # #667eea
                "secondary": PptxRGBColor(118, 75, 162),  # #764ba2
                "text": PptxRGBColor(51, 51, 51),
                "background": PptxRGBColor(255, 255, 255)
            },
            "classic": {
                "primary": PptxRGBColor(0, 51, 102),
                "secondary": PptxRGBColor(204, 153, 0),
                "text": PptxRGBColor(0, 0, 0),
                "background": PptxRGBColor(255, 255, 255)
            },
            "simple": {
                "primary": PptxRGBColor(68, 68, 68),
                "secondary": PptxRGBColor(136, 136, 136),
                "text": PptxRGBColor(34, 34, 34),
                "background": PptxRGBColor(250, 250, 250)
            }
        }

        colors = templates.get(template, templates["modern"])

        # タイトルスライド
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白レイアウト

        # タイトル
        title_box = slide.shapes.add_textbox(
            Inches(1), Inches(2.5), Inches(8), Inches(1.5)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(44)
        title_para.font.bold = True
        title_para.font.color.rgb = colors["primary"]
        title_para.alignment = PP_ALIGN.CENTER

        # サブタイトル
        subtitle_box = slide.shapes.add_textbox(
            Inches(1), Inches(4.5), Inches(8), Inches(0.8)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = f"作成日: {datetime.now().strftime('%Y年%m月%d日')}"
        subtitle_para = subtitle_frame.paragraphs[0]
        subtitle_para.font.size = Pt(18)
        subtitle_para.font.color.rgb = colors["text"]
        subtitle_para.alignment = PP_ALIGN.CENTER

        # 概要スライド
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # タイトル
        heading_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(9), Inches(0.8)
        )
        heading_frame = heading_box.text_frame
        heading_frame.text = "概要"
        heading_para = heading_frame.paragraphs[0]
        heading_para.font.size = Pt(32)
        heading_para.font.bold = True
        heading_para.font.color.rgb = colors["primary"]

        # 概要テキスト
        content_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.5), Inches(9), Inches(5)
        )
        content_frame = content_box.text_frame
        content_frame.word_wrap = True
        content_frame.text = summary
        content_para = content_frame.paragraphs[0]
        content_para.font.size = Pt(18)
        content_para.font.color.rgb = colors["text"]

        # 必要な器具スライド（データにあれば）
        if "required_tools" in data and data["required_tools"]:
            slide = prs.slides.add_slide(prs.slide_layouts[6])

            heading_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.5), Inches(9), Inches(0.8)
            )
            heading_frame = heading_box.text_frame
            heading_frame.text = "必要な器具"
            heading_para = heading_frame.paragraphs[0]
            heading_para.font.size = Pt(32)
            heading_para.font.bold = True
            heading_para.font.color.rgb = colors["primary"]

            tools_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(1.5), Inches(9), Inches(5)
            )
            tools_frame = tools_box.text_frame
            tools_frame.word_wrap = True

            for i, tool in enumerate(data["required_tools"]):
                if i > 0:
                    tools_frame.add_paragraph()
                p = tools_frame.paragraphs[i]
                p.text = f"• {tool}"
                p.font.size = Pt(20)
                p.font.color.rgb = colors["text"]

        # 各ステップのスライド
        steps = data.get("steps", [])
        for step in steps:
            slide = prs.slides.add_slide(prs.slide_layouts[6])

            # ステップタイトル
            step_title = f"ステップ {step.get('step_number', '')}: {step.get('title', '')}"
            heading_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.5), Inches(9), Inches(0.8)
            )
            heading_frame = heading_box.text_frame
            heading_frame.text = step_title
            heading_para = heading_frame.paragraphs[0]
            heading_para.font.size = Pt(28)
            heading_para.font.bold = True
            heading_para.font.color.rgb = colors["primary"]

            # 説明
            desc_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(1.5), Inches(9), Inches(2)
            )
            desc_frame = desc_box.text_frame
            desc_frame.word_wrap = True
            desc_frame.text = step.get('description', '')
            desc_para = desc_frame.paragraphs[0]
            desc_para.font.size = Pt(16)
            desc_para.font.color.rgb = colors["text"]

            # 重要ポイント
            if step.get('key_points'):
                points_box = slide.shapes.add_textbox(
                    Inches(0.5), Inches(4), Inches(4.5), Inches(3)
                )
                points_frame = points_box.text_frame
                points_frame.word_wrap = True

                p = points_frame.paragraphs[0]
                p.text = "重要ポイント:"
                p.font.size = Pt(18)
                p.font.bold = True
                p.font.color.rgb = colors["secondary"]

                for point in step['key_points']:
                    points_frame.add_paragraph()
                    p = points_frame.paragraphs[-1]
                    p.text = f"• {point}"
                    p.font.size = Pt(14)
                    p.font.color.rgb = colors["text"]

            # 注意事項
            if step.get('warnings'):
                warnings_box = slide.shapes.add_textbox(
                    Inches(5), Inches(4), Inches(4.5), Inches(3)
                )
                warnings_frame = warnings_box.text_frame
                warnings_frame.word_wrap = True

                p = warnings_frame.paragraphs[0]
                p.text = "⚠️ 注意事項:"
                p.font.size = Pt(18)
                p.font.bold = True
                p.font.color.rgb = PptxRGBColor(255, 0, 0)

                for warning in step['warnings']:
                    warnings_frame.add_paragraph()
                    p = warnings_frame.paragraphs[-1]
                    p.text = f"• {warning}"
                    p.font.size = Pt(14)
                    p.font.color.rgb = PptxRGBColor(200, 0, 0)

        # ファイルを保存
        file_path = os.path.join(self.output_dir, f"{title}.pptx")
        prs.save(file_path)

        return file_path

    def generate_image_cards(
        self,
        data: Dict,
        title: str,
        template: str = "modern"
    ) -> List[str]:
        """
        画像カードを生成（SNS投稿用）

        Args:
            data: マニュアルデータ
            title: マニュアルタイトル
            template: テンプレート名

        Returns:
            生成した画像ファイルのパスのリスト
        """
        image_paths = []

        # テンプレート設定
        templates = {
            "modern": {
                "bg_color": (102, 126, 234),  # #667eea
                "accent_color": (118, 75, 162),  # #764ba2
                "text_color": (255, 255, 255),
                "card_bg": (255, 255, 255)
            },
            "classic": {
                "bg_color": (0, 51, 102),
                "accent_color": (204, 153, 0),
                "text_color": (255, 255, 255),
                "card_bg": (250, 250, 250)
            }
        }

        colors = templates.get(template, templates["modern"])

        # カードサイズ（Instagram正方形）
        card_size = (1080, 1080)

        # タイトルカード
        img = Image.new('RGB', card_size, color=colors["bg_color"])
        draw = ImageDraw.Draw(img)

        # テキストを描画（簡易版 - フォントサイズは固定）
        try:
            # システムフォントを使用
            title_font = ImageFont.truetype("/System/Library/Fonts/Hiragino Sans GB.ttc", 80)
            body_font = ImageFont.truetype("/System/Library/Fonts/Hiragino Sans GB.ttc", 40)
        except:
            # フォントが見つからない場合はデフォルト
            title_font = ImageFont.load_default()
            body_font = ImageFont.load_default()

        # タイトル
        draw.text(
            (540, 400),
            title,
            fill=colors["text_color"],
            font=title_font,
            anchor="mm"
        )

        # 日付
        date_text = datetime.now().strftime('%Y.%m.%d')
        draw.text(
            (540, 600),
            date_text,
            fill=colors["text_color"],
            font=body_font,
            anchor="mm"
        )

        # 保存
        title_card_path = os.path.join(self.output_dir, f"{title}_card_title.png")
        img.save(title_card_path)
        image_paths.append(title_card_path)

        # 各ステップのカード
        steps = data.get("steps", [])
        for i, step in enumerate(steps[:5]):  # 最大5ステップまで
            img = Image.new('RGB', card_size, color=colors["card_bg"])
            draw = ImageDraw.Draw(img)

            # ヘッダー部分
            draw.rectangle([(0, 0), (1080, 200)], fill=colors["bg_color"])

            # ステップ番号とタイトル
            step_title = f"Step {step.get('step_number', i+1)}"
            draw.text(
                (540, 100),
                step_title,
                fill=colors["text_color"],
                font=title_font,
                anchor="mm"
            )

            # ステップ名
            draw.text(
                (100, 300),
                step.get('title', ''),
                fill=(51, 51, 51),
                font=title_font,
                anchor="lm"
            )

            # 説明（簡略化）
            description = step.get('description', '')[:100] + "..."  # 100文字まで
            y_pos = 450
            for line in self._wrap_text(description, 30):
                draw.text(
                    (100, y_pos),
                    line,
                    fill=(102, 102, 102),
                    font=body_font,
                    anchor="lm"
                )
                y_pos += 50

            # 保存
            step_card_path = os.path.join(
                self.output_dir,
                f"{title}_card_step{i+1}.png"
            )
            img.save(step_card_path)
            image_paths.append(step_card_path)

        return image_paths

    def _wrap_text(self, text: str, width: int) -> List[str]:
        """テキストを指定幅で折り返し"""
        words = text.split()
        lines = []
        current_line = []
        current_length = 0

        for word in words:
            if current_length + len(word) <= width:
                current_line.append(word)
                current_length += len(word) + 1
            else:
                if current_line:
                    lines.append(' '.join(current_line))
                current_line = [word]
                current_length = len(word) + 1

        if current_line:
            lines.append(' '.join(current_line))

        return lines
